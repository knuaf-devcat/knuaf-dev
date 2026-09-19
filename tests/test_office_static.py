"""gg_office.py static checks: relationship screening (H13), legacy school
inspector gating (H14), visible-sheet counting (H15), lineage null guards,
doctor engine logic, workspace-relative receipts and jobs/clean hygiene.

Never launches Word/Excel: only pure functions plus the doctor /
verify-template-lineage / jobs / clean subcommands are exercised.
"""
import hashlib
import importlib.util
import json
import os
import zipfile
from pathlib import Path

import pytest

from conftest import run_script, parse_json

import gg_office

REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OD = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"


def build_docx_zip(path, doc_rels_xml):
    """Minimal OOXML-shaped zip whose word/_rels/document.xml.rels is supplied."""
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr(
            "_rels/.rels",
            f'<Relationships xmlns="{REL_NS}">'
            f'<Relationship Id="rId1" Type="{OD}officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        zf.writestr("word/document.xml", "<document/>")
        zf.writestr("word/_rels/document.xml.rels", doc_rels_xml)
    return path


def rels(*items):
    body = "".join(
        f'<Relationship Id="rId{i}" Type="{t}" Target="{target}" TargetMode="External"/>'
        for i, (t, target) in enumerate(items, 1)
    )
    return f'<Relationships xmlns="{REL_NS}">{body}</Relationships>'


# --------------------------------------------------------------------- H13

def test_attached_template_external_rel_passes(tmp_path):
    z = build_docx_zip(tmp_path / "a.docx", rels((OD + "attachedTemplate", "Normal.dotm")))
    ext = gg_office.check_ooxml_relationships(z)
    assert ext == {("word/_rels/document.xml.rels", OD + "attachedTemplate", "Normal.dotm")}


def test_attached_template_with_file_url_still_passes(tmp_path):
    # Word writes the template as a file URL on some builds; it never fetches content.
    z = build_docx_zip(tmp_path / "a.docx", rels(
        (OD + "attachedTemplate", "file:///Users/x/Library/Group%20Containers/Normal.dotm"),
        (OD + "glossaryDocument", "file:///Users/x/Building%20Blocks.dotx"),
    ))
    gg_office.check_ooxml_relationships(z)


def test_external_link_path_bare_name_passes(tmp_path):
    z = build_docx_zip(tmp_path / "a.docx", rels(
        (OD + "externalLinkPath", "budget.xlsx"),
        (OD + "externalLinkPath", "C:\\Users\\me\\budget.xlsx"),
        (OD + "externalLinkPath", "/Users/me/budget.xlsx"),
    ))
    gg_office.check_ooxml_relationships(z)


def test_file_scheme_hyperlink_rejected(tmp_path):
    z = build_docx_zip(tmp_path / "a.docx", rels((OD + "hyperlink", "file:///etc/passwd")))
    with pytest.raises(ValueError, match="위험한 외부 하이퍼링크"):
        gg_office.check_ooxml_relationships(z)


def test_http_hyperlink_passes_and_unc_ole_rejected(tmp_path):
    gg_office.check_ooxml_relationships(
        build_docx_zip(tmp_path / "ok.docx", rels((OD + "hyperlink", "https://example.org/x")))
    )
    z = build_docx_zip(tmp_path / "bad.docx", rels((OD + "oleObject", "http://evil.example/o.bin")))
    with pytest.raises(ValueError, match="위험한 외부 OOXML 관계 참조"):
        gg_office.check_ooxml_relationships(z)


def test_preexisting_unsafe_rel_tolerated_only_via_baseline(tmp_path):
    key = ("word/_rels/document.xml.rels", OD + "oleObject", "http://evil.example/o.bin")
    z = build_docx_zip(tmp_path / "out.docx", rels((key[1], key[2])))
    with pytest.raises(ValueError):
        gg_office.check_ooxml_relationships(z)
    # Same relationship screened on the input earlier -> not re-rejected on output.
    assert gg_office.check_ooxml_relationships(z, baseline_external={key}) == {key}
    # A different unsafe rel is still fatal even with a baseline.
    z2 = build_docx_zip(tmp_path / "out2.docx", rels((OD + "image", "https://evil.example/i.png")))
    with pytest.raises(ValueError):
        gg_office.check_ooxml_relationships(z2, baseline_external={key})


def test_classify_external_relationship():
    c = gg_office.classify_external_relationship
    assert c(OD + "hyperlink", "mailto:a@b.c")[0]
    assert not c(OD + "hyperlink", "javascript:alert(1)")[0]
    assert c(OD + "attachedTemplate", "smb://host/share/x.dotm")[0]
    assert c(OD + "externalLink", "../other.xlsx")[0]
    assert not c(OD + "externalLink", "file://server/share/other.xlsx")[0]
    assert not c(OD + "image", "https://x/y.png")[0]


# --------------------------------------------------------------------- H14

def test_legacy_school_check_requires_full_ordered_sheet_set():
    pytest.importorskip("openpyxl")
    from gg_school_excel import SCHOOL_SHEETS

    applies = gg_office._legacy_school_check_applies
    assert applies(list(SCHOOL_SHEETS), False)
    assert applies(["anything"], True)
    # partial overlap (a supplied template that happens to share one name) must NOT trigger
    assert not applies([SCHOOL_SHEETS[0], "양식"], False)
    assert not applies(list(SCHOOL_SHEETS[:5]), False)
    # same names, different order -> not the legacy generated workbook
    assert not applies(list(reversed(SCHOOL_SHEETS)), False)
    assert not applies([], False)
    assert applies(("a", "b"), False, school_sheets=("a", "b"))


# --------------------------------------------------------------------- H15

def test_worksheet_count_excludes_hidden_sheets(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    wb.active.title = "Visible"
    wb.create_sheet("Helper").sheet_state = "hidden"
    wb.create_sheet("Secret").sheet_state = "veryHidden"
    wb.create_sheet("Other")
    p = tmp_path / "h.xlsx"
    wb.save(p)
    counts = gg_office.workbook_worksheet_count(p)
    assert counts == {"total": 4, "visible_worksheets": 2, "hidden_worksheets": 2, "chartsheets": 0}


def test_worksheet_count_tells_chartsheets_apart(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.chart import BarChart, Reference
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([1, 2])
    cs = wb.create_chartsheet("Chart")
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=1))
    cs.add_chart(chart)
    p = tmp_path / "c.xlsx"
    wb.save(p)
    counts = gg_office.workbook_worksheet_count(p)
    assert counts["total"] == 2
    assert counts["visible_worksheets"] == 1
    assert counts["chartsheets"] == 1


def test_worksheet_count_no_visible_sheets_raises(tmp_path):
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    p = tmp_path / "allhidden.xlsx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("xl/workbook.xml",
                    f'<workbook xmlns="{ns}" xmlns:r="{r}"><sheets>'
                    f'<sheet name="A" sheetId="1" state="hidden" r:id="rId1"/></sheets></workbook>')
    with pytest.raises(ValueError, match="no visible worksheets"):
        gg_office.workbook_worksheet_count(p)


# ------------------------------------------------- verify-template-lineage

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _clear_receipt(folder):
    src = folder / "source.xlsx"
    out = folder / "cleared.xlsx"
    src.write_bytes(b"source-bytes")
    out.write_bytes(b"cleared-bytes")
    receipt = folder / "clear.json"
    receipt.write_text(json.dumps({
        "schema": "gg-xlsx-template-receipt/v1",
        "source": {"path": str(src), "sha256": sha(src)},
        "output": {"path": str(out), "sha256": sha(out), "status": "blank_template"},
        "cleared": [],
    }), encoding="utf-8")
    return receipt, out


@pytest.mark.parametrize("validation", [
    None,
    {"structure_issues": [], "template_receipt": None},
])
def test_lineage_null_manifest_fields_give_clear_reason(tmp_path, validation):
    clear_receipt, cleared = _clear_receipt(tmp_path)
    native = tmp_path / "native.office.manifest.json"
    native.write_text(json.dumps({
        "status": "converted", "engine": "Microsoft Excel",
        "input_sha256": sha(cleared), "validation": validation,
        "published_office": str(cleared), "published_office_sha256": sha(cleared),
        "published_pdf": str(cleared), "published_pdf_sha256": sha(cleared),
    }), encoding="utf-8")
    res = run_script("gg_office.py", "verify-template-lineage", cleared,
                     "--receipts", clear_receipt, native, "--json")
    assert res.returncode in (1, 2), res.stdout + res.stderr
    out = parse_json(res.stdout)
    assert out["status"] != "pass"
    reason = out.get("reason") or out.get("error") or ""
    assert "NoneType" not in reason
    assert "영수증" in reason


def test_lineage_clear_only_still_passes(tmp_path):
    clear_receipt, cleared = _clear_receipt(tmp_path)
    res = run_script("gg_office.py", "verify-template-lineage", cleared,
                     "--receipts", clear_receipt, "--json")
    assert res.returncode == 0, res.stdout + res.stderr
    assert parse_json(res.stdout)["status"] == "pass"


# ------------------------------------------------------------------ doctor

def _patch_doctor_env(monkeypatch, *, platform, osascript, pypdf, word, excel, pywin32=False):
    monkeypatch.setattr(gg_office.sys, "platform", platform)
    monkeypatch.setattr(gg_office.shutil, "which",
                        lambda name: "/usr/bin/osascript" if (name == "osascript" and osascript) else None)
    real_find_spec = importlib.util.find_spec

    def fake_find_spec(name, *a, **k):
        if name == "pypdf":
            return object() if pypdf else None
        if name == "win32com":
            return object() if pywin32 else None
        return real_find_spec(name, *a, **k)
    monkeypatch.setattr(gg_office.importlib.util, "find_spec", fake_find_spec)
    real_exists = Path.exists

    def fake_exists(self, *a, **k):
        s = str(self)
        if s.endswith("Microsoft Word.app"):
            return word
        if s.endswith("Microsoft Excel.app"):
            return excel
        return real_exists(self, *a, **k)
    monkeypatch.setattr(Path, "exists", fake_exists)
    monkeypatch.setattr(gg_office, "_windows_com_app_registered",
                        lambda prog_id: word if prog_id.startswith("Word") else excel)


def test_doctor_engine_requires_app_and_pypdf(tmp_path, monkeypatch):
    _patch_doctor_env(monkeypatch, platform="darwin", osascript=True, pypdf=True, word=False, excel=False)
    d = gg_office.doctor(workspace=tmp_path / "ws")
    assert d["osascript_available"] is True
    assert d["platform_engine_available"] is True
    assert d["engine_available"] is False
    assert d["word_engine_available"] is False and d["excel_engine_available"] is False

    _patch_doctor_env(monkeypatch, platform="darwin", osascript=True, pypdf=True, word=True, excel=False)
    d = gg_office.doctor(workspace=tmp_path / "ws")
    assert d["word_engine_available"] is True
    assert d["excel_engine_available"] is False
    assert d["engine_available"] is True

    _patch_doctor_env(monkeypatch, platform="darwin", osascript=True, pypdf=False, word=True, excel=True)
    d = gg_office.doctor(workspace=tmp_path / "ws")
    assert d["pypdf_available"] is False
    assert d["engine_available"] is False

    _patch_doctor_env(monkeypatch, platform="darwin", osascript=False, pypdf=True, word=True, excel=True)
    assert gg_office.doctor(workspace=tmp_path / "ws")["engine_available"] is False

    _patch_doctor_env(monkeypatch, platform="win32", osascript=False, pypdf=True, word=False, excel=True, pywin32=True)
    d = gg_office.doctor(workspace=tmp_path / "ws")
    assert d["word_engine_available"] is False and d["excel_engine_available"] is True
    assert d["jobs"] == {"count": 0, "bytes": 0}


def test_doctor_cli_reports_jobs(tmp_path):
    ws = tmp_path / "ws"
    job = ws / "job-0123456789ab-1700000000" / "output"
    job.mkdir(parents=True)
    (job / "x.bin").write_bytes(b"12345")
    (ws / "not-a-job").mkdir()
    res = run_script("gg_office.py", "doctor", "--workspace", ws, "--json")
    assert res.returncode == 0
    out = parse_json(res.stdout)
    assert out["jobs"] == {"count": 1, "bytes": 5}
    for key in ("word_engine_available", "excel_engine_available", "engine_available"):
        assert isinstance(out[key], bool)


# -------------------------------------------------------- receipt paths

def test_receipt_paths_are_workspace_relative(tmp_path):
    ws = tmp_path / "ws"
    job = ws / "job-0123456789ab-1700000000"
    frag = gg_office.job_receipt_paths(
        job, ws, tmp_path / "input" / "보고서.docx",
        job / "output" / "보고서.docx", job / "output" / "보고서.pdf",
    )
    assert frag == {
        "workspace_relative": True,
        "job_id": "job-0123456789ab-1700000000",
        "job_dir": "job-0123456789ab-1700000000",
        "input_file": "보고서.docx",
        "staged_working_file": "job-0123456789ab-1700000000/output/보고서.docx",
        "staged_pdf_file": "job-0123456789ab-1700000000/output/보고서.pdf",
    }
    assert not any(str(tmp_path) in str(v) for v in frag.values())


def test_stage_job_returns_baseline_and_no_absolute_paths_leak(tmp_path):
    ws = tmp_path / "ws"
    src = build_docx_zip(tmp_path / "in.docx", rels((OD + "attachedTemplate", "Normal.dotm")))
    job_dir, staged_in, work, pdf, digest, baseline = gg_office.stage_job(src, ws)
    assert gg_office.JOB_ID_RE.match(job_dir.name)
    assert staged_in.is_file() and work.is_file()
    assert digest == sha(src)
    assert baseline == {("word/_rels/document.xml.rels", OD + "attachedTemplate", "Normal.dotm")}
    listing = gg_office.list_jobs(ws)
    assert listing["count"] == 1 and listing["jobs"][0]["id"] == job_dir.name
    assert listing["jobs"][0]["files"] == 2


# --------------------------------------------------------- jobs / clean

def _make_job(ws, jid, payload=b"abc"):
    d = ws / jid / "output"
    d.mkdir(parents=True)
    (d / "f.bin").write_bytes(payload)
    return ws / jid


def test_jobs_lists_only_well_formed_job_dirs(tmp_path):
    ws = tmp_path / "ws"
    _make_job(ws, "job-0123456789ab-1700000000", b"12345")
    _make_job(ws, "job-ffffffffffff-1700000001", b"1")
    (ws / "job-bad").mkdir()
    (ws / "job-0123456789ab-1700000002").write_text("file, not dir")
    res = run_script("gg_office.py", "jobs", "--workspace", ws, "--json")
    assert res.returncode == 0, res.stdout + res.stderr
    out = parse_json(res.stdout)
    assert out["status"] == "pass"
    assert out["count"] == 2 and out["bytes"] == 6
    ids = [j["id"] for j in out["jobs"]]
    assert ids == ["job-0123456789ab-1700000000", "job-ffffffffffff-1700000001"]
    first = out["jobs"][0]
    assert first["files"] == 1 and first["bytes"] == 5
    assert first["created"] == "2023-11-14T22:13:20Z"


def test_jobs_on_missing_workspace_is_empty(tmp_path):
    res = run_script("gg_office.py", "jobs", "--workspace", tmp_path / "nope", "--json")
    assert res.returncode == 0
    assert parse_json(res.stdout)["jobs"] == []


def test_clean_deletes_only_named_jobs(tmp_path):
    ws = tmp_path / "ws"
    keep = _make_job(ws, "job-aaaaaaaaaaaa-1700000000")
    gone = _make_job(ws, "job-bbbbbbbbbbbb-1700000001", b"xyz")
    res = run_script("gg_office.py", "clean", "--job", gone.name, "--workspace", ws, "--json")
    assert res.returncode == 0, res.stdout + res.stderr
    out = parse_json(res.stdout)
    assert out["status"] == "pass"
    assert out["deleted"] == [{"id": gone.name, "bytes": 3}]
    assert out["refused"] == []
    assert not gone.exists() and keep.exists()


def test_clean_refuses_malformed_outside_and_missing_ids(tmp_path):
    ws = tmp_path / "ws"
    keep = _make_job(ws, "job-aaaaaaaaaaaa-1700000000")
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "victim.txt").write_text("x")
    res = run_script(
        "gg_office.py", "clean",
        "--job", "job-bad",
        "--job", "../outside",
        "--job", "job-aaaaaaaaaaaa-1700000000/output",
        "--job", "job-cccccccccccc-1700000009",
        "--job", str(outside),
        "--workspace", ws, "--json",
    )
    assert res.returncode == 1, res.stdout + res.stderr
    out = parse_json(res.stdout)
    assert out["status"] == "fail"
    assert out["deleted"] == []
    assert len(out["refused"]) == 5
    assert all("작업" in r["reason"] or "거부" in r["reason"] for r in out["refused"])
    assert keep.exists() and (outside / "victim.txt").exists()


def test_clean_refuses_symlinked_job_dir(tmp_path):
    ws = tmp_path / "ws"
    ws.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "victim.txt").write_text("x")
    link = ws / "job-dddddddddddd-1700000000"
    try:
        os.symlink(outside, link, target_is_directory=True)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks unavailable")
    out = gg_office.clean_jobs([link.name], ws)
    assert out["status"] == "fail" and out["deleted"] == []
    assert (outside / "victim.txt").exists()


def test_clean_has_no_all_flag(tmp_path):
    res = run_script("gg_office.py", "clean", "--all", "--workspace", tmp_path, "--json")
    assert res.returncode == 2
    res = run_script("gg_office.py", "clean", "--workspace", tmp_path, "--json")
    assert res.returncode == 2  # --job is required
