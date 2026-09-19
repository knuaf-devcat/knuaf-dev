"""gg_docx_table_layout: keeping numbers off two lines in the school's tables.

A figure broken across lines ("121,24" / "6.7") reads as two numbers. This
module widens columns and steps the font down to prevent that — but it must
never buy a fit with the two things that would matter more: the values
themselves, and a font size nobody can read.
"""
import pytest

pytest.importorskip("docx")

from docx import Document

import gg_docx_table_layout as tl


def table_of(rows):
    doc = Document()
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    return doc, t


FINANCE = [["연도", "매출액", "경비", "순이익"],
           ["2027", "121,246.7", "98,432.1", "22,814.6"],
           ["2028", "133,371.4", "104,287.5", "29,083.9"]]
WIDE = [["항목"] + [f"{i}년차" for i in range(1, 9)],
        ["매출액"] + ["1,234,567.8"] * 8]
TEXT_ONLY = [["구분", "내용"], ["작목", "도라지"], ["지역", "전북 정읍시 북면"]]


# --- what a number is ---------------------------------------------------------

@pytest.mark.parametrize("token", [
    "1,200", "12.5", "121,246.7", "(121,246.7)", "-5", "−5", "–5", "+3",
    "95%", "12.5%", "2026", "0", "(5)", "１２３",
])
def test_figures_a_thesis_actually_prints_are_treated_as_numbers(token):
    assert tl.is_numeric_token(token) is True


@pytest.mark.parametrize("token", ["", "   ", "도라지", "600평", "abc", "...", "–", "%"])
def test_prose_is_not_treated_as_a_number(token):
    assert tl.is_numeric_token(token) is False


def test_a_wider_number_needs_more_room():
    at10 = tl.token_width_mm
    assert at10("121,246.7", 10.0) > at10("1,200", 10.0) > at10("5", 10.0)


def test_a_smaller_font_needs_less_room():
    assert tl.token_width_mm("121,246.7", 9.0) < tl.token_width_mm("121,246.7", 10.0)


# --- the ordinary case --------------------------------------------------------

def test_a_finance_table_fits_at_full_size():
    _, t = table_of(FINANCE)
    report = tl.fit_table(t)
    assert report["applied"] is True
    assert report["feasible"] is True
    assert report["font_pt"] == 10.0, "여유가 있는데 글꼴을 줄임"
    assert report["rotated_to_landscape"] is False
    assert report["warnings"] == []


def test_column_widths_are_assigned_within_the_text_area():
    _, t = table_of(FINANCE)
    report = tl.fit_table(t, text_area_mm=150.0)
    assert len(report["columns"]) == 4
    assert all(w > 0 for w in report["columns"])
    assert sum(report["columns"]) <= 150.0 + 0.5


def test_a_table_with_no_numbers_is_left_alone():
    _, t = table_of(TEXT_ONLY)
    report = tl.fit_table(t)
    assert report["applied"] is False
    assert report["columns"] is None
    assert any("no numeric cells" in w for w in report["warnings"])


# --- the two things it must never trade away ----------------------------------

@pytest.mark.parametrize("rows", [FINANCE, WIDE], ids=["보통표", "과폭표"])
@pytest.mark.parametrize("landscape", [True, False], ids=["가로허용", "가로금지"])
def test_no_cell_text_is_ever_changed(rows, landscape):
    """Layout may move the numbers around the page; it may not edit them."""
    _, t = table_of(rows)
    before = [[c.text for c in row.cells] for row in t.rows]
    tl.fit_table(t, allow_landscape=landscape)
    after = [[c.text for c in row.cells] for row in t.rows]
    assert after == before


@pytest.mark.parametrize("landscape", [True, False], ids=["가로허용", "가로금지"])
def test_the_font_never_drops_below_the_readable_floor(landscape):
    _, t = table_of(WIDE)
    report = tl.fit_table(t, min_font_pt=9.0, allow_landscape=landscape)
    assert report["font_pt"] >= 9.0, "읽을 수 없는 크기로 줄여서 맞춤"


def test_an_impossible_table_reports_failure_instead_of_faking_a_fit():
    """With rotation refused, the honest answer is feasible=False."""
    _, t = table_of(WIDE)
    report = tl.fit_table(t, allow_landscape=False)
    assert report["feasible"] is False
    assert report["font_pt"] == 9.0
    assert report["rotated_to_landscape"] is False
    assert any("cannot fit" in w for w in report["warnings"])
    assert any("values preserved" in w for w in report["warnings"])


def test_a_custom_floor_is_respected_too():
    _, t = table_of(WIDE)
    report = tl.fit_table(t, min_font_pt=9.5, allow_landscape=False)
    assert report["font_pt"] >= 9.5


# --- rotation is opt-in, and it says so --------------------------------------

def test_rotation_only_happens_when_it_is_allowed():
    _, t = table_of(WIDE)
    assert tl.fit_table(t, allow_landscape=False)["rotated_to_landscape"] is False


def test_rotation_when_allowed_is_reported_not_silent():
    _, t = table_of(WIDE)
    report = tl.fit_table(t, allow_landscape=True)
    assert report["rotated_to_landscape"] is True
    assert report["feasible"] is True
    assert any("landscape" in w for w in report["warnings"]), "회전을 말없이 함"
    assert report["text_area_mm"] > 150.0


# --- the report shape build_docx records --------------------------------------

def test_every_report_carries_the_fields_build_docx_writes_to_its_manifest():
    _, t = table_of(FINANCE)
    report = tl.fit_table(t)
    assert set(report) >= {"applied", "feasible", "font_pt", "columns",
                           "warnings", "rotated_to_landscape", "text_area_mm"}


def test_fit_tables_reports_one_entry_per_table():
    doc = Document()
    tables = []
    for rows in (FINANCE, TEXT_ONLY, WIDE):
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
        tables.append(t)
    reports = tl.fit_tables(tables, allow_landscape=False)
    assert len(reports) == 3
    assert [r["applied"] for r in reports] == [True, False, True]
