"""In-process gg_core methods (status, checks, lock, history, restore, export...)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

from .rpc import RpcError

REVIEW_KINDS = ("content", "logic", "calculation", "docx", "xlsx", "render")


def _core():
    import gg_core  # scripts dir is on sys.path

    return gg_core


def _root(params) -> Path:
    root = params.get("root")
    if not isinstance(root, str) or not root:
        raise RpcError("invalid_params", "root(작업 폴더) 필요")
    return Path(root)


def _summary(rows):
    counts = {"pass": 0, "fail": 0, "blocked": 0, "other": 0}
    for r in rows:
        counts[r["status"] if r["status"] in counts else "other"] += 1
    counts["total"] = len(rows)
    return counts


def lanes(checks, gate, completion):
    """Four independent status lanes (SKILL.md invariant 5). Never merged."""
    machine = [r for r in checks if r.get("owner", "skill") == "skill"]
    review_rows = [r for r in gate if r["check_id"].startswith("review_") or r["check_id"] in {"fact_unreviewed", "coverage"}]
    output_rows = [r for r in gate if r["check_id"].startswith("output_")]
    professor = [r for r in gate if r["check_id"] == "professor_approval"]
    by_kind = {}
    for kind in REVIEW_KINDS:
        rows = [r for r in gate if r["check_id"] == "review_" + kind]
        by_kind[kind] = rows[0]["status"] if rows else "pass"
    independent_missing = not any(
        r["check_id"] == "review_" + k and r["status"] == "pass" for r in gate for k in REVIEW_KINDS
    ) and any(v != "pass" for v in by_kind.values())
    return {
        "machine": {"summary": _summary(machine), "ready": bool(completion.get("skill_ready"))},
        "content_review": {"summary": _summary(review_rows), "by_kind": by_kind, "independent_review_missing": independent_missing},
        "output_review": {"summary": _summary(output_rows), "guideline_ready": completion.get("guideline_ready")},
        "professor": {
            "summary": _summary(professor),
            "pending": bool(completion.get("professor_approval_pending")),
            "recorded": any(r["status"] == "pass" for r in professor),
        },
    }


# --- methods -----------------------------------------------------------------


def sys_hello(ctx, params):
    from . import __version__

    return {
        "sidecar_version": __version__,
        "python": sys.executable,
        "python_version": "%d.%d.%d" % sys.version_info[:3],
        "platform": sys.platform,
        "scripts_dir": str(Path(_core().__file__).parent),
        "interpreter_kind": params.get("interpreter_kind") or "unknown",
    }


def project_init(ctx, params):
    return {"project": _core().init(_root(params))}


def project_load(ctx, params):
    p = _core().load(_root(params))
    return {"project": p, "revision": p["revision"], "schema_version": p["schema_version"]}


def project_status(ctx, params):
    core = _core()
    root = _root(params)
    p = core.load(root)
    gate = core.gate(root, p)
    completion = core.completion(root, p, gate)
    checks = core.checks(root, p)
    tasks = core.next_tasks(p)
    return {
        "revision": p["revision"],
        "project_id": p.get("project_id"),
        "checks": checks,
        "gate": gate,
        "completion": completion,
        "tasks": tasks,
        "lanes": lanes(checks, gate, completion),
        "user_finish_pending": completion.get("user_finish_pending", []),
        "counts": {c: len(p[c]) for c in core.COLLECTIONS},
    }


def project_checks(ctx, params):
    core = _core()
    root = _root(params)
    p = core.load(root)
    return core.gate(root, p) if params.get("scope") == "submission" else core.checks(root, p)


def project_next(ctx, params):
    core = _core()
    return core.next_tasks(core.load(_root(params)))


def project_question(ctx, params):
    core = _core()
    field = params.get("field")
    if not field:
        raise RpcError("invalid_params", "field 필요")
    return {"field": field, "action": core.question(core.load(_root(params)), field)}


def project_sections(ctx, params):
    core = _core()
    root = _root(params)
    p = core.load(root)
    out = []
    for sid, s in sorted(p["sections"].items(), key=lambda kv: (kv[1].get("order", 0), kv[0])):
        entry = {k: s.get(k) for k in ("id", "title", "order", "path", "status", "revision", "draft_hash")}
        entry["claims"] = len(s["claims"]) if isinstance(s.get("claims"), list) else None
        try:
            text = core.local(root, s["path"]).read_text(encoding="utf-8")
            entry["draft_hash_ok"] = core.digest(core.draft(text)) == s.get("draft_hash")
            entry["draft_chars"] = len(core.draft(text))
        except (OSError, ValueError) as error:
            entry["draft_hash_ok"] = False
            entry["error"] = str(error)
        out.append(entry)
    return out


def section_read(ctx, params):
    core = _core()
    root = _root(params)
    sid = params.get("id")
    p = core.load(root)
    if sid not in p["sections"]:
        raise RpcError("not_found", "절 없음: %s" % sid)
    s = p["sections"][sid]
    text = core.local(root, s["path"]).read_text(encoding="utf-8")
    return {"id": sid, "title": s.get("title"), "path": s.get("path"), "draft": core.draft(text), "raw_chars": len(text)}


def project_export(ctx, params):
    kind = params.get("kind", "review")
    if kind not in ("draft", "review", "submission_candidate"):
        raise RpcError("invalid_params", "kind는 draft|review|submission_candidate")
    path = _core().export(_root(params), kind)
    return {"path": path, "kind": kind}


def project_history(ctx, params):
    return _core().history(_root(params))


def project_restore(ctx, params):
    rev, expected = params.get("revision"), params.get("expected_revision")
    if type(rev) is not int or type(expected) is not int:
        raise RpcError("invalid_params", "revision, expected_revision 정수 필요")
    return _core().restore(_root(params), rev, expected)


def lock_info(ctx, params):
    return _core().lock_info(_root(params))


def lock_unlock(ctx, params):
    return _core().unlock(_root(params))


def doctor_all(ctx, params):
    core = _core()
    root = _root(params)
    out = {"gg": core.doctor(root), "lock": core.lock_info(root)}
    try:
        import gg_deps

        out["deps"] = gg_deps.doctor(root)
    except Exception as error:  # noqa: BLE001 - diagnostics must not fail the call
        out["deps"] = {"error": str(error)}
    try:
        import gg_kordoc

        cli, kind = gg_kordoc._resolve_kordoc(None)
        cache = gg_kordoc._cache_root(None)
        legacy = [str(r) for r in gg_kordoc._legacy_cache_roots()]
        out["kordoc"] = {
            "cli_path": cli,
            "source_kind": kind,
            "cache_dir": str(cache),
            "cache_present": (cache / gg_kordoc.VERSION).exists() or any((Path(r) / gg_kordoc.VERSION).exists() for r in legacy),
            "legacy_cache_dirs": legacy,
        }
    except Exception as error:  # noqa: BLE001
        out["kordoc"] = {"error": str(error)}
    try:
        out["snapshots"] = core.history(root)
    except (OSError, ValueError) as error:
        out["snapshots"] = {"error": str(error)}
    out["python"] = {"executable": sys.executable, "version": "%d.%d.%d" % sys.version_info[:3]}
    return out


# --- materials / artifacts (read-only views over sources/ and build/) --------

MATERIAL_FIELDS = (
    "intake.current_manuscript",
    "intake.current_finance",
    "intake.reference_materials",
    "intake.work_basis",
)

ARTIFACT_KINDS = {".docx": "docx", ".xlsx": "xlsx", ".pdf": "pdf", ".md": "md"}

XLSX_MAX_ROWS = 200
XLSX_MAX_COLS = 40
TEXT_PREVIEW_LIMIT = 200 * 1024


def _visible_files(base: Path):
    """Files under `base`, skipping hidden names (incl. .knuaf-gui)."""
    if not base.is_dir():
        return
    for f in sorted(base.rglob("*")):
        if not f.is_file():
            continue
        try:
            rel_parts = f.relative_to(base).parts
        except ValueError:
            continue
        if any(part.startswith(".") for part in rel_parts):
            continue
        yield f


def _latest_fact(p, field_id):
    """Newest fact (highest revision) recorded for a field_id, or None."""
    facts = [f for f in (p.get("facts") or {}).values() if f.get("field_id") == field_id]
    if not facts:
        return None
    return max(facts, key=lambda f: (f.get("revision") or 0, str(f.get("id") or "")))


def _fact_path(fact):
    """Where the recorded file lives: `path`, `value.path`, or a path-like string value."""
    path = fact.get("path")
    if isinstance(path, str) and path:
        return path
    value = fact.get("value")
    if isinstance(value, dict):
        inner = value.get("path")
        if isinstance(inner, str) and inner:
            return inner
    if isinstance(value, str) and value:
        return value
    return None


def _fact_entry(fact):
    if fact is None:
        return {"provided": False, "path": None, "value": None, "note": None, "answer_state": None}
    value = fact.get("value")
    return {
        "provided": fact.get("answer_state") == "provided",
        "path": _fact_path(fact),
        "value": value if isinstance(value, str) else None,
        "note": fact.get("note") if isinstance(fact.get("note"), str) else None,
        "answer_state": fact.get("answer_state"),
    }


def _is_current_file(file: Path, current_paths):
    """A sources/ file counts as `current` only when it is the recorded manuscript/finance file itself.
    Same-named files elsewhere stay `reference`: names alone never confirm the current work."""
    return file.resolve() in current_paths["resolved"]


def materials_list(ctx, params):
    """Intake facts + sources/ listing. Display only — nothing is adopted or written here."""
    core = _core()
    root = _root(params).resolve()
    p = core.load(root)
    manuscript = _fact_entry(_latest_fact(p, "intake.current_manuscript"))
    finance = _fact_entry(_latest_fact(p, "intake.current_finance"))
    references = _fact_entry(_latest_fact(p, "intake.reference_materials"))
    basis = _latest_fact(p, "intake.work_basis")
    work_basis = None
    if basis is not None:
        v = basis.get("value")
        work_basis = v if isinstance(v, str) else (basis.get("note") if isinstance(basis.get("note"), str) else None)

    current_paths = {"resolved": set()}
    for entry in (manuscript, finance):
        if not entry["provided"]:
            continue
        raw = entry["path"] or entry["value"]
        if not raw:
            continue
        cand = Path(raw)
        if cand.is_absolute():
            try:
                current_paths["resolved"].add(cand.resolve())
            except OSError:
                pass
        else:
            try:
                current_paths["resolved"].add((root / cand).resolve())
            except OSError:
                pass

    files = []
    sources = root / "sources"
    for f in _visible_files(sources) or ():
        try:
            rel = f.relative_to(root).as_posix()
            st = f.stat()
        except (ValueError, OSError):
            continue
        files.append(
            {
                "path": rel,
                "bytes": st.st_size,
                "mtime": st.st_mtime,
                "role": "current" if _is_current_file(f, current_paths) else "reference",
            }
        )
    return {
        "current_manuscript": manuscript,
        "current_finance": finance,
        "reference_materials": references,
        "work_basis": work_basis,
        "files": files,
    }


def artifact_list(ctx, params):
    """Previewable files under build/, newest first; `revision` from the outputs collection when known."""
    core = _core()
    root = _root(params).resolve()
    p = core.load(root)
    revisions = {}
    for out in (p.get("outputs") or {}).values():
        path = out.get("path")
        if not isinstance(path, str) or not path:
            continue
        key = Path(path).as_posix()
        rev = out.get("revision")
        if isinstance(rev, int) and rev > (revisions.get(key) or 0):
            revisions[key] = rev
    items = []
    for f in _visible_files(root / "build") or ():
        kind = ARTIFACT_KINDS.get(f.suffix.lower())
        if kind is None:
            continue
        try:
            rel = f.relative_to(root).as_posix()
            st = f.stat()
        except (ValueError, OSError):
            continue
        items.append(
            {
                "path": rel,
                "name": f.name,
                "kind": kind,
                "bytes": st.st_size,
                "mtime": st.st_mtime,
                "revision": revisions.get(rel),
            }
        )
    items.sort(key=lambda i: i["mtime"], reverse=True)
    return {"items": items}


def _xlsx_preview(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"kind": "unavailable", "reason": "재무 미리보기에는 프로젝트 패키지가 필요해요. 패키지 준비를 먼저 실행해 주세요."}
    sheets = []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            max_row = ws.max_row or 0
            max_col = ws.max_column or 0
            rows = []
            for row in ws.iter_rows(min_row=1, max_row=min(max_row, XLSX_MAX_ROWS), max_col=min(max_col, XLSX_MAX_COLS)):
                rows.append(["" if cell.value is None else str(cell.value) for cell in row])
            sheets.append(
                {
                    "name": ws.title,
                    "max_row": max_row,
                    "max_col": max_col,
                    "rows": rows,
                    "truncated": max_row > XLSX_MAX_ROWS or max_col > XLSX_MAX_COLS,
                }
            )
    finally:
        wb.close()
    # Second pass with formulas visible: how many formula cells lack a cached value?
    formulas = 0
    uncalculated = 0
    wbf = load_workbook(path, read_only=True, data_only=False)
    try:
        for ws in wbf.worksheets:
            sheet = next((s for s in sheets if s["name"] == ws.title), None)
            max_row = min(ws.max_row or 0, XLSX_MAX_ROWS)
            max_col = min(ws.max_column or 0, XLSX_MAX_COLS)
            for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col)):
                for c_idx, cell in enumerate(row):
                    if cell.data_type != "f":
                        continue
                    formulas += 1
                    if sheet is not None and r_idx < len(sheet["rows"]) and c_idx < len(sheet["rows"][r_idx]) and sheet["rows"][r_idx][c_idx] == "":
                        uncalculated += 1
    finally:
        wbf.close()
    out = {"kind": "xlsx", "sheets": sheets, "formulas": formulas}
    if formulas > 10 and uncalculated / formulas > 0.5:
        out["formulas_uncalculated"] = True
    return out


def artifact_preview(ctx, params):
    """Read-only preview payload; binary kinds only report their kind (bytes go via file:read)."""
    core = _core()
    root = _root(params)
    rel = params.get("path")
    if not isinstance(rel, str) or not rel:
        raise RpcError("invalid_params", "path 필요")
    try:
        f = core.local(root, rel)
    except ValueError as error:
        raise RpcError("invalid_params", str(error))
    if not f.is_file():
        raise RpcError("not_found", "파일 없음: %s" % rel)
    ext = f.suffix.lower()
    if ext == ".xlsx":
        return _xlsx_preview(f)
    if ext in (".md", ".txt"):
        data = f.read_bytes()
        return {
            "kind": "text",
            "text": data[:TEXT_PREVIEW_LIMIT].decode("utf-8", errors="replace"),
            "truncated": len(data) > TEXT_PREVIEW_LIMIT,
        }
    if ext == ".pdf":
        return {"kind": "pdf"}
    return {"kind": "unavailable", "reason": "이 형식은 미리보기를 지원하지 않아요."}


def fs_build_tree(ctx, params):
    core = _core()
    root = _root(params)
    build = root / "build"
    out = {"path": str(build), "revisions": [], "files": []}
    if not build.is_dir():
        return out
    for child in sorted(build.iterdir(), key=lambda p: p.name):
        if child.is_file():
            out["files"].append({"name": child.name, "bytes": child.stat().st_size})
            continue
        if not child.name.isdigit():
            out["files"].append({"name": child.name, "dir": True})
            continue
        rev = {"revision": int(child.name), "kinds": []}
        for kind_dir in sorted(child.iterdir(), key=lambda p: p.name):
            if not kind_dir.is_dir():
                continue
            kind = {"kind": kind_dir.name, "path": str(kind_dir), "files": [], "complete": None}
            try:
                kind["complete"] = core.export_complete(kind_dir, kind_dir.name, rev["revision"])
            except Exception:  # noqa: BLE001
                kind["complete"] = None
            for f in sorted(kind_dir.iterdir(), key=lambda p: p.name):
                if f.is_file() and not f.name.startswith("."):
                    kind["files"].append({"name": f.name, "bytes": f.stat().st_size})
            manifest = kind_dir / "manifest.json"
            if manifest.is_file():
                try:
                    kind["manifest"] = json.loads(manifest.read_text(encoding="utf-8"))
                except (OSError, ValueError):
                    kind["manifest"] = None
            rev["kinds"].append(kind)
        out["revisions"].append(rev)
    return out


def register(server):
    server.register("sys.hello", sys_hello)
    server.register("project.init", project_init, write=True)
    server.register("project.load", project_load)
    server.register("project.status", project_status)
    server.register("project.checks", project_checks)
    server.register("project.next", project_next)
    server.register("project.question", project_question)
    server.register("project.sections", project_sections)
    server.register("section.read", section_read)
    server.register("project.export", project_export, write=True)
    server.register("project.history", project_history)
    server.register("project.restore", project_restore, write=True)
    server.register("lock.info", lock_info)
    server.register("lock.unlock", lock_unlock, write=True)
    server.register("doctor.all", doctor_all)
    server.register("fs.build_tree", fs_build_tree)
    server.register("materials.list", materials_list)
    server.register("artifact.list", artifact_list)
    server.register("artifact.preview", artifact_preview)
