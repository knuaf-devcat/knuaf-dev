"""Measured school form bands. Keeps long titles inside fixed page regions."""
import re
from docx.shared import Pt, Mm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _font(run, font, size, bold=False):
    run.font.name, run.font.size, run.bold = font, Pt(size), bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'): rf.set(qn('w:'+key), font)


def _p(cell, text, font, size=18, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, first=True):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    f = p.paragraph_format
    f.space_before = f.space_after = Pt(0)
    f.line_spacing = Pt(size * 1.35)
    f.keep_with_next = False
    f.widow_control = False
    p.alignment = align
    _font(p.add_run(text), font, size, bold)
    return p


def _cell_props(cell):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for side in ('top','bottom','left','right'):
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'0');e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    borders=OxmlElement('w:tcBorders')
    for side in ('top','bottom','left','right','insideH','insideV'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');borders.append(e)
    pr.append(borders)


def _bands(doc, specs, font):
    # Each start coordinate is millimetres from the physical A4 page top.
    # Section top margin is 20 mm; exact row heights prevent cumulative drift.
    table=doc.add_table(rows=len(specs), cols=1)
    table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width=Mm(150)
    for i,(top,bottom,text,size,bold) in enumerate(specs):
        row=table.rows[i];row.height=Mm(bottom-top);row.height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        cell=row.cells[0];cell.width=Mm(150);_cell_props(cell)
        _p(cell,text,font,size,bold)
    return table


def _break(doc):
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1)
    r=p.add_run();r.font.size=Pt(1)
    br=OxmlElement('w:br');br.set(qn('w:type'),'page');r._r.append(br)


def _texts(nodes,start,end):
    return [n['text'].strip() for n in nodes[start:end] if n.get('text','').strip()]


def render_forms(doc,nodes,plan,font):
    ix=plan['indices']
    cover=_texts(nodes,ix['겉표지']+1,ix['표제면'])
    submission=_texts(nodes,ix['표제면']+1,ix['제출서'])
    approval=_texts(nodes,ix['인준서']+1,ix['목차'])
    if len(cover)<7 or len(submission)<9 or len(approval)<6:
        raise ValueError('학교 양식 필수 메타데이터 순서/항목 누락')
    degree,plan_label=cover[:2]
    graduation,school,department,author=cover[-4:]
    title_parts=cover[2:-4]
    if not 1<=len(title_parts)<=2: raise ValueError('학교 제목과 부제는 각 한 문단으로 입력')
    title=title_parts[0];subtitle=title_parts[1] if len(title_parts)>1 else ''
    import math, unicodedata
    def units(text): return sum(1 if unicodedata.east_asian_width(ch) in {'W','F'} else .55 for ch in text)
    required_height=math.ceil(units(title)/22)*8.573+(math.ceil(units(subtitle)/29)*6.668 if subtitle else 0)
    if required_height>55: raise ValueError('표지 제목/부제가 허용 영역을 초과함: 제목을 줄이거나 별도 양식 프로필 필요')
    for value,limit in ((degree,23),(plan_label,10),(graduation,20),(school,20),(department,22),(author,22)):
        if units(value)>limit: raise ValueError('학교 양식의 한 줄 필드가 너무 김: 이름·학교·학과·날짜 또는 계획명 확인')
    advisor=submission[0]
    if units(advisor)>23: raise ValueError('지도교수 표시가 학교 양식의 한 줄 영역을 초과함')
    submission_date=submission[-4]
    statement=next((s for s in submission if s.startswith('이 논문을')), '이 논문을 농업전문학사 학위논문으로 제출함')
    # 양식 슬롯에 배정되지 않은 저작 행(검토 메모·제출서 문구 등)은 각 양식의
    # 빈 상단 밴드에 그대로 표시해 내용을 조용히 버리지 않는다.
    consumed={advisor,submission_date,statement,degree,plan_label,title,subtitle,graduation,school,department,author}
    extras_sub=[x for x in submission if x not in consumed]
    extras_sub+=_texts(nodes,ix['제출서']+1,ix['인준서'])
    extras_sub=' '.join(extras_sub).strip()
    if extras_sub and units(extras_sub)>85: raise ValueError('표제면/제출서 추가 문구가 양식 상단 영역을 초과함')
    # 인준서 블록은 내용으로 판별한다: (인)으로 끝나는 행은 서명란,
    # '인준함'이 있는 행은 인준 문구, 날짜 행은 연도 패턴, 나머지는 검토 메모.
    signatures=[x for x in approval if x.rstrip().endswith('(인)')]
    appr_core=[x for x in approval if not x.rstrip().endswith('(인)')]
    stmt_i=next((i for i,x in enumerate(appr_core) if '인준함' in x), None)
    if stmt_i is None:
        if len(appr_core)<2: raise ValueError('인준 문구 누락')
        stmt_i=len(appr_core)-2
    appr_statement=appr_core[stmt_i]
    appr_rest=appr_core[:stmt_i]+appr_core[stmt_i+1:]
    date_i=next((i for i,x in enumerate(appr_rest) if re.search(r'\d{4}\s*년',x)), None)
    appr_date=appr_rest.pop(date_i) if date_i is not None else (appr_rest.pop() if appr_rest else '')
    if len(signatures)>4: raise ValueError('인준 서명란은 최대 4명')
    appr_extras=' '.join(appr_rest).strip()
    if appr_extras and units(appr_extras)>140: raise ValueError('인준서 추가 문구가 양식 상단 영역을 초과함')
    table=_bands(doc,[
        (20,52,'',1,False),(52,88,degree,18,False),(88,116,' '.join(plan_label) if len(plan_label)<=9 else plan_label,30,False),
        (116,176,title,18,False),(176,218,graduation,20,False),(218,233,school,20,False),
        (233,248,department,18,False),(248,265,author,18,False)],font)
    if subtitle: _p(table.cell(3,0),subtitle,font,14,first=False)
    _break(doc)
    table=_bands(doc,[
        (20,40,extras_sub,12,False),(40,51,advisor,18,False),(51,88,degree,18,False),
        (88,116,' '.join(plan_label) if len(plan_label)<=9 else plan_label,30,False),(116,175,title,18,False),
        (175,188,statement,16,False),(188,224,submission_date,20,False),(224,238,school,20,False),
        (238,252,department,18,False),(252,268,author,18,False)],font)
    if subtitle: _p(table.cell(4,0),subtitle,font,14,first=False)
    _break(doc)
    table=_bands(doc,[(20,52,appr_extras,12,False),(52,95,appr_statement,18,False),(95,159,appr_date,20,False),
                      (159,185,'',1,False),(185,211,'',1,False),(211,237,'',1,False),(237,266,'',1,False)],font)
    for i,line in enumerate(signatures):
        cell=table.cell(i+3,0);p=cell.paragraphs[0]
        p.paragraph_format.line_spacing=Pt(1);p.runs[0].font.size=Pt(1)
        inner=cell.add_table(rows=1,cols=3);inner.autofit=False;inner.alignment=WD_TABLE_ALIGNMENT.RIGHT
        widths=[24,65,19]
        label='위원장' if i==0 else '위  원'
        name=line.removeprefix('위원장').removeprefix('위  원').removesuffix('(인)').strip()
        for j,(w,text) in enumerate(zip(widths,[label,name,'(인)'])):
            inner.columns[j].width=Mm(w);c=inner.cell(0,j);c.width=Mm(w);_cell_props(c)
            _p(c,text,font,18,align=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER)
            borders=c._tc.get_or_add_tcPr().find(qn('w:tcBorders'));e=borders.find(qn('w:bottom'));e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000')
        inner.rows[0].height=Mm(9);inner.rows[0].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
        tail=cell.paragraphs[-1];tail.paragraph_format.line_spacing=Pt(1);tail.paragraph_format.space_after=Pt(0)
    _break(doc)


def bookmark(p,name,num):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(num));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(num))
    p._p.append(start);p._p.append(end)


def page_reference(p,name,font,size=12):
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGEREF '+name+' \\h');field.set(qn('w:dirty'),'true')
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(size*2));pr.append(sz);run.append(pr)
    txt=OxmlElement('w:t');txt.text='?';run.append(txt);field.append(run);p._p.append(field)


def contents(doc,entries,font,title='목    차',anchor=None,num=9000):
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Mm(18);p.paragraph_format.space_after=Mm(12)
    _font(p.add_run(title),font,20)
    if anchor: bookmark(p,anchor,num)
    for label,name,level in entries:
        p=doc.add_paragraph();f=p.paragraph_format
        f.space_before=f.space_after=Pt(0);f.line_spacing=Pt(21)
        f.left_indent=Mm(5 if level==2 else 0)
        f.first_line_indent=Mm(0)
        f.tab_stops.add_tab_stop(Mm(150), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        _font(p.add_run(label),font,12,level==1)
        p.add_run('\t');page_reference(p,name,font)
